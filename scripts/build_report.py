"""Builds the LifeChain mini-project report (DOCX) following the format of the reference lab document."""
from pathlib import Path
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
SHOTS = ROOT / "screenshots"
OUT = ROOT / "report" / "LifeChain_MiniProject_Report.docx"
OUT.parent.mkdir(exist_ok=True)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)  # A4
sec.left_margin = sec.right_margin = Inches(0.9)
sec.top_margin, sec.bottom_margin = Inches(0.75), Inches(0.75)
TEXT_W = 8.27 - 1.8

st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(10)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
st.paragraph_format.space_after = Pt(5)
st.paragraph_format.line_spacing = 1.0

fig_no = 0
J = WD_ALIGN_PARAGRAPH.JUSTIFY


def para(text="", bold_prefix=None, align=J, space_after=None, keep=False):
    p = doc.add_paragraph()
    p.alignment = align
    if bold_prefix:
        p.add_run(bold_prefix).bold = True
    p.add_run(text)
    if space_after is not None:
        p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.keep_with_next = keep
    return p


def heading(text, size=14, center=False, before=8):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(size)
    return p


def h1(text): return heading(text, 14)
def h2(text): return heading(text, 12.5, before=6)
def h3(text): return heading(text, 11, before=4)


def page_heading(text):
    """Large centered heading at the top of a new page (as on pages 2-4 of the reference)."""
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.paragraphs[-1].paragraph_format.space_after = Pt(0)
    return heading(text, 18, center=True, before=0)


def bullets(items):
    for it in items:
        p = doc.add_paragraph(f"•  {it}")
        p.alignment = J
        p.paragraph_format.space_after = Pt(3)


def numbered(items):
    for i, it in enumerate(items, 1):
        p = doc.add_paragraph(f"{i}. {it}")
        p.alignment = J
        p.paragraph_format.space_after = Pt(3)


def figure(name, caption, width=TEXT_W):
    global fig_no
    fig_no += 1
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(3)
    p.add_run().add_picture(str(SHOTS / name), width=Inches(width))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.add_run(f"Figure {fig_no}: {caption}").font.size = Pt(9)
    c.paragraph_format.space_after = Pt(10)


def table(headers, rows, widths, font=9.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(font)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].paragraphs[0].add_run(str(v)).font.size = Pt(font)
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Inches(w)
            pf = row.cells[i].paragraphs[0].paragraph_format
            pf.space_before = Pt(3)
            pf.space_after = Pt(3)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


# ======================================================================= PAGE 1
p = doc.add_paragraph()
p.paragraph_format.tab_stops.add_tab_stop(Inches(TEXT_W / 2), WD_TAB_ALIGNMENT.CENTER)
p.paragraph_format.tab_stops.add_tab_stop(Inches(TEXT_W), WD_TAB_ALIGNMENT.RIGHT)
p.add_run("Name : Vishal Bhor\tId : VU1F2324041\tDiv : B / A").font.size = Pt(11)
p.paragraph_format.space_after = Pt(16)

t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("MINI PROJECT")
r.bold = True
r.underline = True
r.font.size = Pt(16)
t.paragraph_format.space_after = Pt(10)

para("To design and implement LifeChain, a permissioned blockchain application that tracks every blood unit from donor to "
     "patient, and to understand how a blockchain network can make blood bank management safe, transparent and auditable.",
     bold_prefix="Aim: ")

h1("1. Introduction to LifeChain")
para("Blood cannot be manufactured; it can only come from donors. Every blood unit passes through several organizations before it "
     "reaches a patient: a blood bank collects it, a laboratory screens it, and a hospital transfuses it. Each step is usually recorded in "
     "a separate system, so the exact status, safety and location of a unit is difficult to know at any moment.")
para("LifeChain is a permissioned blockchain network whose known members are a blood bank, a testing laboratory, hospitals and a health "
     "regulator. Each organization records its actions on a blood unit as signed transactions, which are endorsed, ordered and validated "
     "before they update a shared, tamper-evident ledger.")

h2("1.1 Purpose and Key Features")
bullets([
    "Permissioned membership: participants use identities issued by their organization's certificate authority.",
    "Shared ledger: all organizations maintain one consistent record of every blood unit.",
    "Policy enforcement: every transaction must be endorsed by the submitting organization and the regulator.",
    "Safety rules: untested, unsafe, expired or incompatible blood cannot be issued or transfused.",
    "Privacy: donor and patient identities stay off-chain; only pseudonymous references are stored.",
])

h1("2. LifeChain Network Architecture")
para("The LifeChain network contains five organizations, one peer per organization, an ordering service, identity services and a client "
     "application. Peers host the chaincode and ledger data. The ordering service arranges endorsed transactions into blocks, which are "
     "sent to peers for validation and commitment. The Membership Service Provider (MSP) identifies trusted certificates and membership.")
para("The ledger includes a blockchain that preserves the history of every blood-unit transaction and a world state that stores the latest "
     "status of each unit. Chaincode contains the business logic, while the endorsement policy defines which peers must approve a transaction.")
figure("fig1_architecture.png", "LifeChain permissioned blockchain network architecture", width=5.3)

h2("2.1 Main Components")
for label, text in [
    ("Organizations", "blood bank, testing laboratory, two hospitals and the health regulator, each managing its own users and peer."),
    ("Peers", "nodes that simulate proposals, endorse transactions, validate blocks and maintain ledger data."),
    ("Ordering service", "establishes a common order of transactions and packages them into hash-linked blocks."),
    ("CA and MSP", "certificates identify users and nodes; the MSP determines which identities are trusted."),
    ("Chaincode (bloodtrace)", "smart-contract logic that defines valid operations on blood units and who may perform them."),
]:
    para(f"{label}: {text}", space_after=4)

# ======================================================================= PAGE 2
page_heading("3. Transaction Processing in LifeChain")
para("LifeChain follows the execute–order–validate model of permissioned blockchains such as Hyperledger Fabric. The endorsing peers "
     "simulate a proposed change before the network orders it.")
numbered([
    "Proposal: a user signs a request (for example CollectBlood) with the identity of their organization and sends it to the peers required by the endorsement policy.",
    "Execution / simulation: the peers run the bloodtrace chaincode on their world state and prepare the read set and write set without changing the ledger.",
    "Endorsement: each peer signs the result with its ECDSA key. The gateway checks that the required endorsements (creator organization + regulator) have been collected.",
    "Ordering: the endorsed transaction is sent to the ordering service, which sequences transactions and creates a block containing the hash of the previous block.",
    "Distribution: the block is delivered to all peers on the channel bloodchannel.",
    "Validation: peers check the endorsement signatures and confirm that the versions of the keys read during simulation have not changed (MVCC check).",
    "Commit: the block is appended to the ledger. Valid transactions update the world state; invalid transactions remain in the block but do not apply their changes.",
    "Result: the application shows the transaction ID, block number and validation code to the user.",
])
h2("3.1 Blockchain and World State")
para("The blockchain keeps the chronological record of blocks and transactions, which supports audit and recall investigations. The world "
     "state is a current snapshot of every blood unit. For example, the blockchain shows each recorded event of unit BU-2026-0113 "
     "(collected, tested, dispatched, received, transfused), while the world state simply shows its latest status: TRANSFUSED.")
h2("3.2 Endorsement, Ordering and Validation")
para("The endorsement policy of LifeChain is AND('<CreatorOrg>.peer', 'RegulatorMSP.peer'): every change must be approved by the peer of the "
     "submitting organization and by the regulator's peer. The ordering service then sequences endorsed transactions, but ordering alone "
     "does not make a transaction valid. Each peer validates the transaction before applying it. If two hospitals are allocated the "
     "same blood unit at the same time, only the first transaction is valid and the second is rejected with MVCC_READ_CONFLICT.")

h1("4. Case Study: Blood Bank Traceability")
para("This scenario shows how LifeChain follows a blood unit from the donor to the patient. The goal is to provide authorized "
     "participants with a common, auditable history of collection, screening, storage, distribution and transfusion of every unit.")
h3("4.1 Problem")
para("Blood bank records are often kept in separate registers, spreadsheets or isolated software at each blood bank, laboratory and "
     "hospital. As a result, hospitals cannot see real-time stock by blood group, units sometimes expire unused, and records can be "
     "edited without trace. When a donor is later found to be infected or a patient has a transfusion reaction, tracing the affected units "
     "and the patients who received them takes a long time, because records must be collected and compared across organizations.")
h3("4.2 Proposed Blockchain-Based Approach")
para("A permissioned blockchain network records each blood-unit event as it happens. Each approved organization submits transactions only "
     "for the actions it is responsible for, and the chaincode enforces safety rules such as mandatory testing, expiry checks and blood-group "
     "compatibility. Access control and pseudonymous references protect donor and patient privacy.")
h3("4.3 Participants and Responsibilities")
table(["Participant", "Responsibilities"], [
    ["Blood Bank (LifeLine Blood Bank)", "Registers donations with unit ID, blood group, component, volume and expiry; stores tested units and issues them to hospitals."],
    ["Testing Laboratory (Precision Diagnostics Lab)", "Records screening results for HIV, HBV, HCV, Syphilis and Malaria; marks each unit SAFE or UNSAFE."],
    ["Hospitals (CityCare, Green Valley)", "Confirm receipt of issued units and record transfusions to patients after compatibility checks."],
    ["Regulator (Health Regulatory Authority)", "Endorses every transaction and uses read-only access to audit stock, provenance and event history."],
    ["Donor / Patient", "Not network members; represented only by pseudonymous references (REF-…) to protect personal data."],
], widths=[2.2, TEXT_W - 2.2])

# ======================================================================= PAGE 3
page_heading("4.4 Data Model and Workflow")
para("A blood-unit record contains a unique unit ID, pseudonymous donor reference, blood group, component (whole blood, packed RBC, "
     "platelets or plasma), volume, collection time and centre, expiry date, test results, current status, current custodian and "
     "transfusion details. Personal details of donors and patients are kept off-chain; only a SHA-256 based reference is stored.")
h2("Possible Chaincode Operations")
bullets([
    "CollectBlood: register a new blood unit; expiry is calculated from the component's shelf life.",
    "RecordTest: add screening results; the unit becomes AVAILABLE only if all five tests are negative.",
    "TransferUnit: issue a tested, non-expired unit from its current custodian to a registered hospital.",
    "ReceiveUnit: the destination hospital confirms receipt and becomes the new custodian.",
    "TransfuseUnit: record a transfusion after checking expiry and ABO/Rh compatibility.",
    "DiscardUnit: record disposal of expired or unsafe units.",
    "QueryUnitHistory: retrieve the complete history of a unit from the blockchain.",
])
h2("Workflow")
numbered([
    "The blood bank signs in with an approved identity and registers a donated unit.",
    "The client sends the proposal to the peers required by the endorsement policy.",
    "The peers simulate the chaincode and return signed endorsements.",
    "The endorsed transaction is sent to the ordering service and included in a block.",
    "The peers validate the block and commit valid changes to the ledger and world state.",
    "The laboratory records test results, the blood bank issues the unit, and the hospital receives and transfuses it.",
    "If a problem occurs, the regulator queries the unit history to identify its donor, test report and recipient.",
])
figure("fig2_unit_movement.png", "Blood unit movement and traceability events", width=6.0)

h2("4.5 Expected Benefits")
bullets([
    "A consistent event history helps authorized users trace every blood unit from donor to patient.",
    "Real-time stock by blood group and expiry helps hospitals find blood quickly and reduces wastage of expired units.",
    "Untested, unsafe, expired or incompatible blood is blocked automatically by the chaincode.",
    "Transactions are linked to identities and endorsement policies, which improves accountability.",
    "Recall and adverse-reaction investigations become faster because records are shared across organizations.",
])
h2("4.6 Practical Limitations")
bullets([
    "The ledger cannot prove that the physical blood bag matches the data entered; barcode/RFID labelling is still required.",
    "Participants must agree on governance, membership, endorsement rules and dispute handling.",
    "Patient and donor privacy laws require careful design of what is stored on-chain.",
    "Integration with existing blood bank software, laboratory instruments and hospital systems needs extra development.",
    "Infrastructure, maintenance, staff training and performance must be evaluated before deployment.",
])

# ======================================================================= IMPLEMENTATION
page_heading("5. Implementation and Output")
para("The LifeChain prototype implements the network described above. The ledger, chaincode, endorsement, ordering and validation logic "
     "are written in Node.js using its built-in crypto library, and a web application allows each organization to perform its role. "
     "The identity selector in the application acts as the user's wallet: every request is signed with the selected organization's identity.")
h2("5.1 Technology Used")
table(["Component", "Technology"], [
    ["Ledger and chaincode", "Node.js (JavaScript): blocks, SHA-256 hashing, Merkle data hash, world state with key versions"],
    ["Identities and signatures", "ECDSA P-256 key pairs for each organization (simulated CA / MSP)"],
    ["Consensus", "Endorsement policy + ordering service + MVCC validation (execute–order–validate)"],
    ["Back-end", "Node.js HTTP server exposing a REST API (client gateway)"],
    ["Front-end", "HTML5, CSS3 and JavaScript single-page web application"],
    ["Storage", "data/ledger.json (block store); world state rebuilt by replaying valid transactions"],
], widths=[2.0, TEXT_W - 2.0])
h2("5.2 Steps to Run the Application")
numbered([
    "Install Node.js (version 18 or later) and open a terminal in the project folder.",
    "Run npm run seed to create the network and sample blood units (optional).",
    "Run npm start to start the peer, ordering service and gateway.",
    "Open http://localhost:3000 in a browser and choose an organization identity from the sidebar.",
])
h2("5.3 Output Screenshots")
SHOTS_LIST = [
    ("00_terminal.png", "Creating the network with sample data and starting the LifeChain gateway",
     "Twelve donations are registered in block 1 and screened in block 2 (one unit is rejected as HBV positive); later blocks record transfers, receipts, transfusions and a discard. The ledger is validated at start-up.", 5.7),
    ("01_dashboard.png", "Dashboard: ledger status, stock by blood group and latest blocks",
     "Available stock is shown per blood group (out-of-stock groups in red, low stock in yellow) along with expiring units, hash-linked blocks and recent transactions.", None),
    ("02_collect_form.png", "Blood bank registering a new donation (unit BU-2026-0113, B+ whole blood)",
     "Signed in as LifeLine Blood Bank, the user enters the unit barcode, donor ID, blood group, component and volume.", None),
    ("03_collect_committed.png", "CollectBlood transaction endorsed by BloodBankMSP and RegulatorMSP and committed in block #13",
     "The receipt shows the transaction ID, block hash, both peer endorsements, the pseudonymous donor reference and the calculated expiry date.", None),
    ("04_access_denied.png", "Access control: a hospital is not allowed to register blood donations",
     "When CityCare Hospital tries to invoke CollectBlood, the endorsing peers reject the proposal, so it is never ordered into a block.", None),
    ("05_lab_test_recorded.png", "Testing laboratory recording screening results (block #14)",
     "All five infection tests are negative, so the chaincode marks the unit SAFE and its status changes to AVAILABLE.", None),
    ("06_distribution_issued.png", "Blood bank issuing the tested unit to CityCare Hospital (block #15)",
     "Only tested, non-expired units are listed for issue; expiry is shown with the number of days remaining.", None),
    ("07_hospital_received.png", "CityCare Hospital confirming receipt of the unit (block #16)",
     "Only the hospital to which the unit was addressed can receive it; it then becomes the unit's custodian.", None),
    ("08_incompatible_blocked.png", "Incompatible transfusion blocked by the chaincode",
     "An attempt to give B+ whole blood to an A+ patient is rejected by the ABO/Rh compatibility rule in the chaincode.", None),
    ("09_transfusion_recorded.png", "Compatible transfusion recorded (block #17)",
     "The B+ unit is transfused to a B+ patient. The patient is stored only as a pseudonymous reference (REF-…).", None),
    ("10_inventory.png", "Inventory: world state of all blood units",
     "Every unit with its blood group, component, donor reference, expiry, current custodian and status, as seen by the regulator.", None),
    ("11_trace_timeline.png", "Donor-to-patient provenance of unit BU-2026-0113",
     "The trace query lists all five ledger events (collected, tested safe, dispatched, received, transfused) with block numbers, organizations and endorsements.", None),
    ("12_block_explorer.png", "Block explorer showing hash-linked blocks",
     "Each block shows its hash, previous-block hash and data hash, and each transaction shows its validation code and endorsing organizations.", None),
    ("13_network.png", "Network members, channel configuration and access-control rules",
     "Five organizations, each with its own MSP, peer, certificate authority and key fingerprint, on channel bloodchannel.", None),
    ("14_double_allocation.png", "Double allocation prevented by MVCC validation (block #18)",
     "Unit BU-2026-0106 is issued to two hospitals at the same time. Both proposals are endorsed, but only the first is VALID; the second fails with MVCC_READ_CONFLICT.", None),
    ("15_tamper_detected.png", "Tamper detection: blood group changed directly in a stored block",
     "Changing the blood group of BU-2026-0101 from O+ to O- in block #1 breaks the block's data hash and both endorsement signatures, and the block is highlighted in red.", None),
]
for img, cap, desc, w in SHOTS_LIST:
    para(desc, space_after=3, keep=True)
    figure(img, cap, width=w or 5.9)

# ======================================================================= OTHER USES ...
page_heading("6. Other Uses of Blockchain in Healthcare")
for label, text in [
    ("Organ donation", "hospitals and transplant authorities can maintain a transparent waiting list and matching record."),
    ("Pharmaceutical supply chain", "manufacturers, distributors and pharmacies can trace medicines and detect counterfeit drugs."),
    ("Vaccine cold chain", "temperature and custody events can be recorded to prove that vaccines were stored correctly."),
    ("Medical records", "patients can grant hospitals permission to access references to their health records."),
    ("Health insurance", "insurers and hospitals can track claim progress and supporting documents."),
    ("Clinical trials", "trial data and consent records can be time-stamped to prevent later alteration."),
]:
    para(f"{label}: {text}", space_after=4)

h1("7. Advantages")
bullets([
    "Known identities and controlled membership support trusted collaboration between blood banks, laboratories and hospitals.",
    "Endorsement policies ensure that the regulator approves every change to the ledger.",
    "Chaincode applies safety rules (testing, expiry, compatibility) consistently in every organization.",
    "The ledger maintains a complete transaction history together with the latest status of every unit.",
    "MVCC validation prevents the same blood unit from being allocated twice.",
    "Any modification of stored blocks is detected through hash links, data hashes and signatures.",
])
h1("8. Challenges and Limitations")
bullets([
    "Network setup and maintenance are more involved than a single blood bank database.",
    "Organizations need shared governance and agreed operating procedures.",
    "Security depends on proper certificate handling, key protection, backups and monitoring.",
    "Incorrect data entered at collection can still be recorded on an immutable ledger.",
    "The prototype runs all peers on one machine; a real deployment needs separate nodes for each organization.",
    "A blockchain is not necessary when one trusted organization controls all the data.",
])
h1("9. Security and Privacy Measures")
bullets([
    "Use trusted identities and protect the private keys and certificates of every organization.",
    "Grant only the permissions needed for each role (e.g. only laboratories may record test results).",
    "Require regulator endorsement for every transaction.",
    "Store only pseudonymous donor and patient references on the ledger; keep personal data off-chain.",
    "Monitor peer and ordering-service health, keep backups and plan certificate renewal.",
    "Validate inputs in the chaincode; immutability alone does not guarantee data accuracy.",
])
h1("10. Conclusion")
para("LifeChain shows how a permissioned blockchain can make blood bank management safer and more transparent. Its separation of "
     "execution, ordering, validation and commitment, with identities and endorsement policies, ensures that only authorized organizations "
     "record events and that all share the same history. In the application, every unit can be traced from donor to patient, unsafe, expired "
     "and incompatible blood is blocked, double allocation is prevented and tampering is detected. A practical deployment still needs "
     "accurate data entry, physical labelling, sound governance, privacy controls and integration with existing systems.")
h1("11. References")
numbered([
    "Blockchain Technology Laboratory Manual, Experiment 9: Hyperledger Fabric Architecture and Application.",
    "Hyperledger Fabric Documentation: https://hyperledger-fabric.readthedocs.io/",
    "Hyperledger Foundation — Fabric project information: https://www.hyperledger.org/use/fabric",
    "World Health Organization, Blood safety and availability: https://www.who.int/news-room/fact-sheets/detail/blood-safety-and-availability",
    "Node.js Crypto Module Documentation: https://nodejs.org/api/crypto.html",
])

doc.save(OUT)
print("saved", OUT)
